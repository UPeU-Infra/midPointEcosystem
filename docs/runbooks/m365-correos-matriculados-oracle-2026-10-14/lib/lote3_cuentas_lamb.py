#!/usr/bin/env python3
"""
Lote 3 de CORREO_INST: alumnos de la hoja «Sin cuenta» que SÍ tienen cuenta en M365 — la creada con su usuario
LAMB (ELISEO.USERS, por ID_PERSONA), casi siempre en la migración Gmail→Outlook de jun-2022, cuyo displayName es
el propio usuario y por eso el cruce por nombre no la encontró.

  python3 lote3_cuentas_lamb.py <dir_trabajo> <Correos_crear_y_licenciar_*.xlsx>

<dir_trabajo> debe tener: entra_hoy.json (pull del tenant), eliseo_users.tsv (ID, EMAIL de ELISEO.USERS),
ora_correo_inst_hoy.tsv (ID_PERSONA, ID_TIPODOCUMENTO, NUM_DOCUMENTO, CORREO_INST_RAW, CORREO_INST — leído HOY),
ora_doc_dup.tsv (fichas con documento duplicado).
Salida en ~/Downloads: Excel del lote 3, UPDATE (sin COMMIT, candados por ficha/documento/valor actual) y ROLLBACK,
y un Excel con los casos que necesitan revisión.
"""
import csv, json, os, re, sys, unicodedata, datetime
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

W, XL = sys.argv[1], sys.argv[2]
OUT = os.path.expanduser("~/Downloads"); FECHA = datetime.date.today().isoformat()
def s(x): return re.sub(r"[^a-z]", "", unicodedata.normalize("NFKD", (x or "").lower()).encode("ascii", "ignore").decode())
def lp(a): return (a or "").split("@")[0].lower()

idx = {}
for u in json.load(open(f"{W}/entra_hoy.json")):
    for a in [u["upn"], u["mail"]] + [p.split(":", 1)[1] for p in (u["proxy"] or [])]:
        if a and a.lower().endswith("@upeu.edu.pe"): idx.setdefault(lp(a), u)
lamb = {}
for line in list(open(f"{W}/eliseo_users.tsv", encoding="utf-8", errors="replace"))[1:]:
    p = line.rstrip("\n").split("\t")
    if len(p) > 1 and p[1].strip(): lamb[p[0].strip()] = lp(p[1].strip())
actual, dueno = {}, defaultdict(set)
for r in csv.DictReader(open(f"{W}/ora_correo_inst_hoy.tsv", encoding="utf-8", errors="replace"), delimiter="\t"):
    actual[r["ID_PERSONA"].strip()] = r["CORREO_INST_RAW"]; dueno[r["CORREO_INST"].strip()].add(r["ID_PERSONA"].strip())
docdup = defaultdict(set)
for r in csv.DictReader(open(f"{W}/ora_doc_dup.tsv", encoding="utf-8", errors="replace"), delimiter="\t"):
    docdup[(r["ID_TIPODOCUMENTO"], r["NUM_DOCUMENTO"])].add(r["ID_PERSONA"])

rows = list(openpyxl.load_workbook(XL, read_only=True)["Sin cuenta"].iter_rows(values_only=True))
h = rows[0]; R = [dict(zip(h, ["" if v is None else str(v).strip() for v in r])) for r in rows[1:]]
cargar, revisar, crear, uso = [], [], [], Counter()
for r in R:
    k = r["ID_PERSONA"]; u_l = lamb.get(k, ""); acc = idx.get(u_l) if u_l else None
    if not acc: crear.append(r); continue
    correo = (acc.get("mail") or acc["upn"]).lower()
    blob = s(acc["upn"]) + s(acc["display"])
    nombres = [s(t) for t in r["NOMBRES"].split() if len(s(t)) >= 3]
    pat = s(r["PATERNO"].split()[0]) if r["PATERNO"] else ""
    nombre_ok = (any(t in blob for t in nombres) and pat and pat in blob)
    cur = (actual.get(k) or "").strip().lower()
    motivo = ""
    if not nombre_ok: motivo = f"el usuario LAMB «{u_l}» existe pero su nombre en M365 («{acc['display']}») no coincide claramente"
    elif not acc.get("enabled"): motivo = "la cuenta está deshabilitada en M365"
    elif not acc.get("mail"): motivo = "la cuenta no tiene buzón (falta licencia)"
    elif dueno.get(correo, set()) - {k}: motivo = f"ese correo ya es el CORREO_INST de otra ficha (ID_PERSONA {', '.join(sorted(dueno[correo] - {k}))})"
    elif docdup.get((r["ID_TIPODOCUMENTO"], r["NUM_DOCUMENTO"]), set()) - {k}: motivo = "su documento tiene otra ficha en Oracle"
    elif cur and cur.endswith("@upeu.edu.pe") and not lp(cur).isdigit() and cur != correo: motivo = f"Oracle ya tiene «{cur}»"
    elif cur == correo: motivo = "Oracle ya tiene este correo (nada que cargar)"
    uso[correo] += 1
    (revisar if motivo else cargar).append((r, acc, correo, cur, motivo))
dobles = {c for c, n in uso.items() if n > 1}
cargar2 = [x for x in cargar if x[2] not in dobles]
revisar += [(r, a, c, cur, "el mismo correo resultó para 2 personas") for r, a, c, cur, _ in cargar if c in dobles]
cargar = cargar2

def hoja(wb, nombre, cab, filas):
    ws = wb.create_sheet(nombre); ws.append(cab)
    for c in ws[1]: c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79"), Alignment(wrap_text=True)
    for f in filas: ws.append(f)
    for i in range(1, len(cab) + 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 24
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
base = lambda r: [int(r["ID_PERSONA"]), r["ID_TIPODOCUMENTO"], r["NUM_DOCUMENTO"], r["CODIGO"], r["PATERNO"], r["MATERNO"], r["NOMBRES"], r["SEDE"], r["PROGRAMA"]]
CAB = ["ID_PERSONA", "ID_TIPODOCUMENTO", "NUM_DOCUMENTO", "CODIGO", "PATERNO", "MATERNO", "NOMBRES", "SEDE", "PROGRAMA"]
wb = openpyxl.Workbook(); wb.remove(wb.active)
hoja(wb, "Cargar en Oracle", CAB + ["CORREO_INST_ACTUAL", "CORREO_A_CARGAR", "CUENTA_CREADA"],
     [base(r) + [cur, c, (a.get("created") or "")[:10]] for r, a, c, cur, _ in sorted(cargar, key=lambda x: (x[0]["SEDE"], x[0]["PATERNO"]))])
x1 = f"{OUT}/Correos_para_cargar_Oracle_pregrado_2026-2_lote3_{FECHA}.xlsx"; wb.save(x1)
wb = openpyxl.Workbook(); wb.remove(wb.active)
hoja(wb, "Revisar", CAB + ["CUENTA_M365", "NOMBRE_EN_M365", "MOTIVO", "RESULTADO"],
     [base(r) + [c, a.get("display"), m, ""] for r, a, c, cur, m in sorted(revisar, key=lambda x: (x[4], x[0]["PATERNO"]))])
x2 = f"{OUT}/Correos_lote3_por_revisar_{FECHA}.xlsx"; wb.save(x2)

q = lambda v: "'" + str(v).replace("'", "''") + "'"
up = [f"-- Carga de CORREO_INST · matriculados de pregrado 2026-2 · lote 3 (cuentas LAMB ya existentes en M365) · {FECHA}",
      f"-- Filas: {len(cargar)}. Candados: ID_PERSONA + tipo y nº de documento + CORREO_INST igual al leído hoy. Sin COMMIT.",
      "SET SERVEROUTPUT ON SIZE UNLIMITED", "SET DEFINE OFF", ""]
rb = [f"-- ROLLBACK del lote 3 ({FECHA}). Sin COMMIT.", "SET SERVEROUTPUT ON SIZE UNLIMITED", "SET DEFINE OFF", ""]
for i in range(0, len(cargar), 400):
    lot = cargar[i:i + 400]
    for dst, inv in ((up, False), (rb, True)):
        dst.append("DECLARE n NUMBER := 0; BEGIN")
        for r, a, c, cur, _ in lot:
            raw = actual.get(r["ID_PERSONA"])
            if not inv:
                dst.append(f"  UPDATE MOISES.PERSONA_NATURAL SET CORREO_INST = {q(c)} WHERE ID_PERSONA = {int(r['ID_PERSONA'])}"
                           f" AND ID_TIPODOCUMENTO = {int(r['ID_TIPODOCUMENTO'])} AND TRIM(NUM_DOCUMENTO) = {q(r["NUM_DOCUMENTO"])}"
                           f" AND NVL(LOWER(TRIM(CORREO_INST)),'#') = {q(cur or '#')}; n := n + SQL%ROWCOUNT;")
            else:
                dst.append(f"  UPDATE MOISES.PERSONA_NATURAL SET CORREO_INST = {'NULL' if not raw else q(raw)} WHERE ID_PERSONA = {int(r['ID_PERSONA'])}"
                           f" AND LOWER(TRIM(CORREO_INST)) = {q(c)}; n := n + SQL%ROWCOUNT;")
        dst.append(f"  DBMS_OUTPUT.PUT_LINE('bloque {i // 400 + 1}: ' || n || ' de {len(lot)} filas');"); dst += ["END;", "/", ""]
partes = ["(ID_PERSONA, LOWER(TRIM(CORREO_INST))) IN (" + ",".join(f"({int(r['ID_PERSONA'])},{q(c)})" for r, a, c, cur, _ in cargar[i:i + 500]) + ")"
          for i in range(0, len(cargar), 500)]
up += [f"-- VERIFICACIÓN: debe devolver {len(cargar)}",
       "SELECT COUNT(*) AS FILAS_CON_CORREO_ESPERADO FROM MOISES.PERSONA_NATURAL WHERE\n   " + "\n OR ".join(partes) + ";",
       "", f"-- Si FILAS_CON_CORREO_ESPERADO = {len(cargar)} → COMMIT;   si no → ROLLBACK;"]
sq = f"{OUT}/UPDATE_CORREO_INST_pregrado_2026-2_lote3_{FECHA}.sql"; open(sq, "w").write("\n".join(up) + "\n")
rq = f"{OUT}/ROLLBACK_CORREO_INST_pregrado_2026-2_lote3_{FECHA}.sql"; open(rq, "w").write("\n".join(rb) + "\n")
print(f"de {len(R)} «sin cuenta»: lote 3 a cargar {len(cargar)} · revisar {len(revisar)} · crear de verdad {len(crear)}")
print("motivos de revisión:", dict(Counter(re.sub(r"«.*»|\(.*\)", "", m).strip() for *_, m in revisar)))
print(x1); print(x2); print(sq); print(rq)
