#!/usr/bin/env python3
"""
Elimina las cuentas que crear_cuentas.py creó para alumnos que YA tenían cuenta (de la migración Gmail→Outlook de
jun-2022, con displayName = usuario, por eso no se detectó) y apunta Oracle a la cuenta que el alumno ya usa.
Decisión de Alberto (15-oct): «si tienen duplicados, borra el duplicado siempre y cuando esa cuenta no ha sido usada nunca».

  python3 deduplicar_cuentas.py simular                       # SOLO LECTURA
  python3 deduplicar_cuentas.py ejecutar <plan_dedup_*.json> [--limite N]

Cuenta previa = DUPLICADO SEGURO si Oracle la tiene como CORREO (personal) de esa misma ficha; MUY PROBABLE si su
dirección es nombre.paterno del alumno, no es de nadie más en Oracle (CORREO_INST, CORREO, usuario LAMB, employeeId), nadie
más en Oracle se llama igual (nombre + paterno) y su nombre visible es el propio usuario (migración 2022) o lleva
los DOS apellidos del alumno. El resto con cuenta parecida es DUDOSO → Excel para Mesa de
Servicios, no se toca.
ejecutar: por cada alumno vuelve a comprobar que la cuenta que creamos NO tiene ningún intento de inicio de sesión
(auditLogs/signIns) y que la previa sigue habilitada; entonces DELETE de la nuestra (papelera 30 días: restaurable con
POST /directory/deletedItems/{id}/restore). Genera el SQL de CORREO_INST con la cuenta previa (candado = valor leído hoy).
"""
import json, os, re, sys, glob, csv, time, datetime, subprocess, urllib.parse
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import crear_cuentas as cc      # utilidades (token, graph, slug); su main no corre al importar
log, graph, slug, tokens, lp, HOME, OUTD, DOM = cc.log, cc.graph, cc.slug, cc.tokens, cc.lp, cc.HOME, cc.OUTD, cc.DOM

env = {}
for p in ("midpoint-upeu.env", "oracle-lamb.env"): env.update(cc.envfile(os.path.join(HOME, ".secrets", p)))
JAR = "/opt/midpoint/var/lib/ojdbc11-23.6.0.24.10.jar"; DSN = "jdbc:oracle:thin:@192.168.13.9:1521/UPEU"
def ora(q):
    cmd = f"docker exec midpoint_server java -cp /tmp:{JAR} OraQ '{DSN}' '{env['ORACLE_USER']}' '{env['ORACLE_PASS']}' \"{q}\""
    r = subprocess.run(["sshpass", "-e", "ssh", "-o", "StrictHostKeyChecking=no", "-o", "ConnectTimeout=60", "-o", "NumberOfPasswordPrompts=1",
                        "midpoint-prod", cmd], env={**os.environ, "SSHPASS": env["MIDPOINT_PROD_PASS"]}, capture_output=True, text=True, timeout=900)
    if r.returncode: sys.exit(f"✗ Oracle: {r.stderr[-300:]}")
    return [l.split("\t") for l in r.stdout.splitlines()[1:] if l.strip()]
N = "TRANSLATE(UPPER(NOMBRE),'ÁÉÍÓÚÑÜ','AEIOUNU')"; P = "TRANSLATE(UPPER(PATERNO),'ÁÉÍÓÚÑÜ','AEIOUNU')"

def creadas():
    C = {}
    for f in sorted(glob.glob(os.path.join(HOME, "backups/entra-snapshots/creacion_resultado_*.json"))):
        for x in json.load(open(f))["hechos"]:
            if x.get("id"): C[x["ID_PERSONA"]] = x
    return C

def intentos(uid):
    """Intentos de inicio de sesión de una cuenta (cualquier resultado). None si no se pudo consultar."""
    st, j = graph("GET", "/auditLogs/signIns?$top=5&$filter=" + urllib.parse.quote(f"userId eq '{uid}'"))
    return len(j["value"]) if st == 200 else None

def simular():
    C = creadas()
    log(f"cuentas creadas por el IGA: {len(C)} · leyendo el tenant (~3 min)…")
    T, url = [], "/users?$select=id,userPrincipalName,displayName,employeeId,createdDateTime,accountEnabled&$top=999"
    while url:
        st, j = graph("GET", url)
        if st != 200: sys.exit(f"✗ /users {st}")
        T += j["value"]; url = j.get("@odata.nextLink")
    nuestros = {x["id"] for x in C.values()}
    vivas = {x["id"] for x in T}
    C = {k: x for k, x in C.items() if x["id"] in vivas}          # las ya eliminadas no cuentan
    pers, inst, lamb = {}, {}, {}
    for k, c in ora("SELECT ID_PERSONA, LOWER(TRIM(CORREO)) FROM MOISES.PERSONA_NATURAL WHERE LOWER(CORREO) LIKE '%@upeu.edu.pe'"):
        pers.setdefault(lp(c), set()).add(k)
    for rr in csv.DictReader(open(os.path.join(HOME, ".cache/upeu-correos-oracle/ora_correo_inst_hoy.tsv"), encoding="utf-8", errors="replace"), delimiter="\t"):
        c = (rr.get("CORREO_INST") or "").strip().lower()
        if c.endswith("@" + DOM): inst.setdefault(lp(c), set()).add(rr["ID_PERSONA"].strip())
    for idp, u in cc.usuarios_lamb().items(): lamb.setdefault(u, set()).add(idp)
    dueno = lambda l: pers.get(l, set()) | inst.get(l, set()) | lamb.get(l, set())
    seguro, prob, dudoso = [], [], []
    filas = []
    for k, x in C.items():
        mio = {x["upn"].lower()}
        # 1) SEGURO: su correo personal en Oracle es otra cuenta @upeu viva
        previa = [u for u in T if u["id"] not in nuestros and u["userPrincipalName"].lower().endswith("@" + DOM)
                  and k in pers.get(lp(u["userPrincipalName"]), set())]
        if previa:
            seguro.append((k, x, previa[0], "Oracle la tiene como su correo personal")); continue
        giv = [slug(t) for t in tokens(x["NOMBRES"]) if len(slug(t)) >= 3]; pat = slug(x["PATERNO"])
        cands = []
        for u in T:
            l = lp(u["userPrincipalName"])
            if u["id"] in nuestros or not u["userPrincipalName"].lower().endswith("@" + DOM) or l.isdigit(): continue
            b = slug(l) + "|" + slug(u.get("displayName") or "")
            if pat not in b or not any(g in b for g in giv): continue
            if dueno(l) - {k} or (u.get("employeeId") and u["employeeId"] != k): continue
            # el nombre visible no puede ser de OTRA persona: o es el propio usuario (migración 2022, sin datos)
            # o contiene los dos apellidos del alumno («fiorella.cruz» = «Fiorella Victoria De La Cruz Cuyutupac» ≠ Nayeli Fiorella Cruz Chambi)
            d = slug(u.get("displayName") or "")
            mat = slug(x["MATERNO"] or "")
            if not (d == slug(l) or (pat in d and mat and mat in d)): continue
            g = next((g for g in giv if g in b), giv[0])
            cands.append((u, g.upper(), pat.upper()))
        if cands: filas.append((k, x, cands))
    pares = sorted({(g, p) for _, _, cs in filas for _, g, p in cs}); cnt = {}
    for i in range(0, len(pares), 60):
        q = " UNION ALL ".join(f"SELECT '{g}' G, '{p}' P, COUNT(*) C FROM MOISES.PERSONA WHERE {P}='{p}' AND ' '||{N}||' ' LIKE '% {g} %'" for g, p in pares[i:i + 60])
        for g, p, c in ora(q): cnt[(g, p)] = int(c)
    for k, x, cs in filas:
        unicas = [u for u, g, p in cs if cnt.get((g, p)) == 1]
        if len(unicas) == 1: prob.append((k, x, unicas[0], "nombre.paterno sin otro dueño en Oracle y sin homónimos"))
        else: dudoso.append((k, x, [u for u, _, _ in cs]))
    plan = []
    for cls, L in (("SEGURO", seguro), ("MUY PROBABLE", prob)):
        for k, x, u, ev in L:
            plan.append({"clase": cls, "evidencia": ev, "ID_PERSONA": k, "CODIGO": x["CODIGO"], "ID_TIPODOCUMENTO": x["ID_TIPODOCUMENTO"],
                         "NUM_DOCUMENTO": x["NUM_DOCUMENTO"], "NOMBRES": x["NOMBRES"], "PATERNO": x["PATERNO"], "MATERNO": x["MATERNO"],
                         "SEDE": x["SEDE"], "nuestra_id": x["id"], "nuestra_upn": x["upn"].lower(),
                         "previa_id": u["id"], "previa_upn": u["userPrincipalName"].lower(), "previa_nombre": u.get("displayName"),
                         "previa_creada": (u.get("createdDateTime") or "")[:10], "previa_habilitada": u.get("accountEnabled")})
    # una cuenta previa no puede quedar para dos alumnos
    from collections import Counter
    rep = {p for p, n in Counter(q["previa_upn"] for q in plan).items() if n > 1}
    for q in [q for q in plan if q["previa_upn"] in rep]:
        plan.remove(q); dudoso.append((q["ID_PERSONA"], C[q["ID_PERSONA"]], []))
    fecha = datetime.date.today().isoformat()
    pj = os.path.join(HOME, f"backups/entra-snapshots/plan_dedup_{fecha}.json")
    json.dump(plan, open(pj, "w"), ensure_ascii=False, indent=1); os.chmod(pj, 0o600)
    import openpyxl as ox
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = ox.Workbook(); ws = wb.active; ws.title = "Se eliminaría la nueva"
    ws.append(["CLASE", "ID_PERSONA", "CODIGO", "NOMBRES", "PATERNO", "MATERNO", "SEDE", "CUENTA_QUE_YA_TENÍA", "NOMBRE_EN_M365", "CREADA", "CUENTA_NUEVA_A_ELIMINAR", "EVIDENCIA"])
    for q in plan: ws.append([q["clase"], q["ID_PERSONA"], q["CODIGO"], q["NOMBRES"], q["PATERNO"], q["MATERNO"], q["SEDE"], q["previa_upn"], q["previa_nombre"], q["previa_creada"], q["nuestra_upn"], q["evidencia"]])
    xo = os.path.join(OUTD, f"Simulacion_deduplicar_cuentas_{fecha}.xlsx"); wb.save(xo)
    wb = ox.Workbook(); ws = wb.active; ws.title = "Confirmar cuenta"
    ws.append(["ID_PERSONA", "CODIGO", "NOMBRES", "PATERNO", "MATERNO", "SEDE", "CUENTA_NUEVA (creada 15-oct)", "CUENTA_ANTERIOR_POSIBLE", "NOMBRE_EN_M365", "RESULTADO (es suya / no es suya)"])
    for k, x, us in dudoso:
        for u in (us or [{}]):
            ws.append([k, x["CODIGO"], x["NOMBRES"], x["PATERNO"], x["MATERNO"], x["SEDE"], x["upn"].lower(), (u.get("userPrincipalName") or "").lower(), u.get("displayName"), ""])
    for c in ws[1]: c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79"), Alignment(wrap_text=True)
    for i in range(1, ws.max_column + 1): ws.column_dimensions[ox.utils.get_column_letter(i)].width = 24
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    xf = os.path.join(OUTD, f"Cuentas_duplicadas_por_confirmar_Flor_{fecha}.xlsx"); wb.save(xf)
    log(f"\nSE ELIMINARÍA LA NUEVA: {len(plan)} {dict(Counter(q['clase'] for q in plan))} · previa deshabilitada: {sum(not q['previa_habilitada'] for q in plan)}")
    log(f"DUDOSOS (a Mesa de Servicios): {len(dudoso)} alumnos")
    log(f"Excel: {xo}\nFlor: {xf}\nPlan: {pj}")

def ejecutar(pj, limite):
    plan = json.load(open(pj))[: limite or None]
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
    hechos, err = [], []
    for i, q in enumerate(plan, 1):
        st, u = graph("GET", f"/users/{q['nuestra_id']}?$select=id,userPrincipalName,employeeId")
        if st == 404: log(f"  ya eliminada {q['nuestra_upn']}"); hechos.append({**q, "estado": "ya eliminada"}); continue
        if st != 200 or str(u.get("employeeId")) != q["ID_PERSONA"]:
            err.append((q["ID_PERSONA"], q["nuestra_upn"], f"lectura nuestra {st}")); continue
        st, v = graph("GET", f"/users/{q['previa_id']}?$select=id,userPrincipalName,accountEnabled")
        if st != 200 or not v.get("accountEnabled"):
            err.append((q["ID_PERSONA"], q["previa_upn"], f"cuenta previa no disponible ({st})")); log(f"  ✗ previa {q['previa_upn']} no disponible → no se toca"); continue
        n = intentos(q["nuestra_id"])
        if n is None or n > 0:
            err.append((q["ID_PERSONA"], q["nuestra_upn"], f"intentos de inicio de sesión: {n}")); log(f"  ✗ {q['nuestra_upn']} tiene uso ({n}) → no se borra"); continue
        st, r = graph("DELETE", f"/users/{q['nuestra_id']}")
        if st != 204: err.append((q["ID_PERSONA"], q["nuestra_upn"], f"DELETE {st} {str(r)[:120]}")); log(f"  ✗ DELETE {q['nuestra_upn']}: {st}"); continue
        log(f"  ✓ eliminada {u['userPrincipalName'].lower()} · se queda con {v['userPrincipalName'].lower()}  ({i}/{len(plan)})")
        hechos.append({**q, "nuestra_upn": u["userPrincipalName"].lower(), "previa_upn": v["userPrincipalName"].lower(), "estado": "eliminada"})
        time.sleep(0.5)
    out = os.path.join(HOME, f"backups/entra-snapshots/dedup_resultado_{stamp}.json")
    json.dump({"hechos": hechos, "errores": err}, open(out, "w"), ensure_ascii=False, indent=1); os.chmod(out, 0o600)
    ok = [h for h in hechos if h["estado"] in ("eliminada", "ya eliminada")]
    if ok:
        actual = {r[0]: r[1] for r in ora("SELECT ID_PERSONA, NVL(LOWER(TRIM(CORREO_INST)),'#') FROM MOISES.PERSONA_NATURAL WHERE ID_PERSONA IN (" + ",".join(h["ID_PERSONA"] for h in ok) + ")")}
        qq = lambda s: "'" + str(s).replace("'", "''") + "'"
        filas = [h for h in ok if actual.get(h["ID_PERSONA"], "#") != h["previa_upn"]]
        sql = [f"-- CORREO_INST = la cuenta que el alumno YA usaba (se eliminó el duplicado creado el 15-oct) · {stamp} · {len(filas)} filas · sin COMMIT",
               "-- Candados: ID_PERSONA + tipo y nº de documento + CORREO_INST igual al leído hoy.", "SET SERVEROUTPUT ON SIZE UNLIMITED", "SET DEFINE OFF", "DECLARE n NUMBER := 0; BEGIN"]
        rb = [f"-- ROLLBACK ({stamp}). Sin COMMIT.", "SET SERVEROUTPUT ON SIZE UNLIMITED", "SET DEFINE OFF", "DECLARE n NUMBER := 0; BEGIN"]
        for h in filas:
            cur = actual.get(h["ID_PERSONA"], "#")
            sql.append(f"  UPDATE MOISES.PERSONA_NATURAL SET CORREO_INST = {qq(h['previa_upn'])} WHERE ID_PERSONA = {int(h['ID_PERSONA'])}"
                       f" AND ID_TIPODOCUMENTO = {int(h['ID_TIPODOCUMENTO'])} AND TRIM(NUM_DOCUMENTO) = {qq(h['NUM_DOCUMENTO'])}"
                       f" AND NVL(LOWER(TRIM(CORREO_INST)),'#') = {qq(cur)}; n := n + SQL%ROWCOUNT;")
            rb.append(f"  UPDATE MOISES.PERSONA_NATURAL SET CORREO_INST = {'NULL' if cur == '#' else qq(cur)} WHERE ID_PERSONA = {int(h['ID_PERSONA'])}"
                      f" AND LOWER(TRIM(CORREO_INST)) = {qq(h['previa_upn'])}; n := n + SQL%ROWCOUNT;")
        for d in (sql, rb): d += [f"  DBMS_OUTPUT.PUT_LINE('actualizadas: ' || n || ' de {len(filas)}');", "END;", "/"]
        sql.append(f"-- Si 'actualizadas' = {len(filas)} → COMMIT; si no → ROLLBACK;")
        for nm, d in (("UPDATE", sql), ("ROLLBACK", rb)):
            o = os.path.join(OUTD, f"{nm}_CORREO_INST_cuentas_duplicadas_resueltas_{stamp}.sql"); open(o, "w").write("\n".join(d) + "\n"); os.chmod(o, 0o600); log(o)
    log(f"\nELIMINADAS: {sum(h['estado']=='eliminada' for h in hechos)} · ya estaban: {sum(h['estado']=='ya eliminada' for h in hechos)} · errores: {len(err)}")
    log(f"resultado: {out}")

if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in ("simular", "ejecutar"): sys.exit(__doc__)
    if sys.argv[1] == "simular": simular()
    else:
        lim = int(sys.argv[sys.argv.index("--limite") + 1]) if "--limite" in sys.argv else 0
        ejecutar(sys.argv[2], lim)
