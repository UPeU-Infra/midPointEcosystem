#!/usr/bin/env python3
"""
Crea en M365 las cuentas de los matriculados de pregrado que no tienen ninguna, con la nomenclatura oficial
UPeU («Nueva Sintaxis de Correo Electrónico de Estudiantes»), y entrega el Excel + SQL para Oracle.

  python3 crear_cuentas.py simular  <Correos_crear_y_licenciar_*.xlsx>            # SOLO LECTURA
  python3 crear_cuentas.py ejecutar <plan_creacion_*.json> [--limite N]          # escribe en M365

Regla de nombre (documento oficial + práctica descrita por Owen en la reunión del 14-oct):
  primernombre.apellidopaterno@upeu.edu.pe, minúsculas, sin tildes/ñ/espacios/apóstrofos;
  si ya existe → + 2 últimos dígitos del código; si vuelve a existir → + 3 últimos dígitos.
Si LAMB ya le asignó un usuario al alumno (ELISEO.USERS, por ID_PERSONA) y es válido y está libre en M365,
SE USA ESE (así M365, LAMB y Oracle coinciden); la regla solo se aplica si no hay usuario LAMB utilizable.
«Ocupado» = cualquier UPN, mail o proxyAddress (alias) del tenant, cuentas BORRADAS incluidas (retienen sus
direcciones 30 días), el usuario LAMB de OTRA persona y lo ya propuesto en el mismo lote.

Quedan RETENIDOS (no se crean) y van a una hoja aparte:
  · apellido paterno compuesto (regla pendiente de validar en el anexo);
  · con posible cuenta existente (POSIBLE_CUENTA_EN_M365): la revisa Mesa de Servicios;
  · Oracle ya trae una dirección @upeu que no existe en M365 (decisión del Ing. Urquizo).

Ejecución: contraseña inicial vigente de admisión desde ~/.secrets/lamb-cuentas.env (LAMB_INITIAL_PASSWORD),
con forceChangePasswordNextSignIn + forceChangePasswordNextSignInWithMfa; usageLocation PE; employeeId =
ID_PERSONA (ancla de correlación); licencia A3 estudiante con reintentos. Antes de crear cada cuenta vuelve
a comprobar EN VIVO que la dirección está libre. Idempotente: si la cuenta ya existe con ese employeeId, la salta.
"""
import json, os, re, sys, time, datetime, unicodedata, urllib.request, urllib.parse, urllib.error

def log(*a): print(*a, flush=True)
HOME = os.path.expanduser("~")
OUTD = os.path.join(HOME, "Downloads")
A3 = "18250162-5d87-4436-a834-d795c15c80f3"   # M365EDU_A3_STUUSEBNFT
DOM = "upeu.edu.pe"

def envfile(p):
    d = {}
    if os.path.exists(p):
        for line in open(p):
            line = line.strip()
            if "=" in line and not line.startswith("#"):
                k, v = line.split("=", 1); d[k.replace("export ", "").strip()] = v.strip().strip('"').strip("'")
    return d
E = envfile(os.path.join(HOME, ".secrets/upeu-infra.env"))
TOK = json.load(urllib.request.urlopen(urllib.request.Request(
    f"https://login.microsoftonline.com/{E['MIDPOINT_AZ_TENANT_ID']}/oauth2/v2.0/token",
    data=urllib.parse.urlencode({"client_id": E["MIDPOINT_AZ_CLIENT_ID"], "client_secret": E["MIDPOINT_AZ_CLIENT_SECRET"],
                                 "scope": "https://graph.microsoft.com/.default", "grant_type": "client_credentials"}).encode()),
    timeout=30))["access_token"]

def graph(method, url, body=None, eventual=False, intentos=5):
    h = {"Authorization": "Bearer " + TOK, "Content-Type": "application/json"}
    if eventual: h["ConsistencyLevel"] = "eventual"
    full = url if url.startswith("http") else "https://graph.microsoft.com/v1.0" + url
    for n in range(intentos):
        try:
            r = urllib.request.urlopen(urllib.request.Request(full, method=method, headers=h,
                                       data=json.dumps(body).encode() if body is not None else None), timeout=30)
            raw = r.read(); return r.status, (json.loads(raw) if raw else None)
        except urllib.error.HTTPError as e:
            txt = e.read().decode(errors="replace")
            if e.code in (429, 500, 502, 503, 504) and n < intentos - 1:
                time.sleep(int(e.headers.get("Retry-After") or 5 * (n + 1))); continue
            return e.code, txt
        except Exception as e:
            if n < intentos - 1: time.sleep(5 * (n + 1)); continue
            return 0, f"{type(e).__name__}: {e}"

def slug(s):
    s = unicodedata.normalize("NFKD", (s or "").lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z]", "", s)
def tokens(s): return [t for t in re.split(r"\s+", (s or "").strip()) if t]
def lp(a): return (a or "").split("@")[0].lower()

def usuarios_lamb():
    """ELISEO.USERS: ID (= ID_PERSONA) → username que LAMB asignó al alumno."""
    eu, d = os.path.join(HOME, ".cache/upeu-correos-oracle/eliseo_users.tsv"), {}
    if os.path.exists(eu):
        for line in list(open(eu, encoding="utf-8", errors="replace"))[1:]:
            p = line.rstrip("\n").split("\t")
            if len(p) > 1 and p[1].strip(): d[p[0].strip()] = lp(p[1].strip())
    return d

def direcciones_ocupadas():
    """Local-parts ocupados en el tenant M365 (activos y borrados: retienen sus direcciones 30 días)."""
    occ = set()
    for base in ("/users", "/directory/deletedItems/microsoft.graph.user"):
        url = f"{base}?$select=userPrincipalName,mail,proxyAddresses&$top=999"; n = 0
        while url:
            st, j = graph("GET", url)
            if st != 200: sys.exit(f"✗ no se pudo leer {base}: {st} {str(j)[:200]}")
            for u in j["value"]:
                for a in [u.get("userPrincipalName"), u.get("mail")] + [p.split(":", 1)[1] for p in (u.get("proxyAddresses") or [])]:
                    if a and "@" in a: occ.add(lp(a))
            n += len(j["value"]); url = j.get("@odata.nextLink")
        log(f"  {base}: {n} objetos leídos")
    return occ

# ───────────────────────────── SIMULAR ─────────────────────────────
def simular(xlsx):
    import openpyxl
    rows = list(openpyxl.load_workbook(xlsx, read_only=True)["Sin cuenta"].iter_rows(values_only=True))
    h = rows[0]; R = [dict(zip(h, ["" if v is None else str(v).strip() for v in r])) for r in rows[1:]]
    log(f"alumnos sin cuenta: {len(R)} · leyendo direcciones ocupadas del tenant (~3 min)…")
    occ = direcciones_ocupadas()
    lamb = usuarios_lamb()
    if not lamb: sys.exit("✗ falta ~/.cache/upeu-correos-oracle/eliseo_users.tsv (usuarios de LAMB)")
    lamb_dueno = {}
    for idp, u in lamb.items(): lamb_dueno.setdefault(u, set()).add(idp)
    log(f"  ELISEO.USERS: {len(lamb)} usuarios LAMB")
    VALIDO = re.compile(r"^[a-z0-9]+(\.[a-z0-9]+)*$")
    plan, ret = [], []
    for r in R:
        n, p = tokens(r["NOMBRES"]), tokens(r["PATERNO"])
        cur = (r["CORREO_INST_ACTUAL_ORACLE"] or "").lower()
        propio0 = lamb.get(r["ID_PERSONA"], "")
        if propio0 and propio0 in occ:
            ret.append((r, f"YA TIENE CUENTA: su usuario LAMB «{propio0}» existe en M365 → va en el lote 3 de Oracle, no se crea")); continue
        if r["POSIBLE_CUENTA_EN_M365"]:
            ret.append((r, f"posible cuenta existente: {r['POSIBLE_CUENTA_EN_M365']} ({r['NOMBRE_EN_M365']})")); continue
        if cur.endswith("@" + DOM) and not lp(cur).isdigit():
            ret.append((r, f"Oracle ya trae {cur}, que no existe en M365: decisión pendiente")); continue
        if not n or not p:
            ret.append((r, "sin nombre o apellido paterno en Oracle")); continue
        if len(p) > 1:
            ret.append((r, f"apellido paterno compuesto «{r['PATERNO']}»: regla pendiente del anexo")); continue
        base = f"{slug(n[0])}.{slug(p[0])}"
        cod = re.sub(r"\D", "", r["CODIGO"])
        propio = lamb.get(r["ID_PERSONA"], "")
        ajeno = lambda c: bool(lamb_dueno.get(c, set()) - {r["ID_PERSONA"]})   # usuario LAMB de OTRA persona
        elegido = regla = None
        if propio and VALIDO.match(propio) and propio not in occ and not ajeno(propio) and not propio.isdigit():
            elegido, regla = propio, "usuario LAMB del alumno"
        else:
            for cand, rg in ((base, "regla: base"), (base + cod[-2:], "regla: +2 dígitos"), (base + cod[-3:], "regla: +3 dígitos")):
                if cand not in occ and not ajeno(cand):
                    elegido, regla = cand, rg; break
        if not elegido:
            ret.append((r, f"{base}, +2 y +3 dígitos ocupados: requiere decisión manual")); continue
        occ.add(elegido)   # reservar dentro del lote
        nombre_vis = " ".join(x.capitalize() for x in tokens(r["NOMBRES"]) + tokens(r["PATERNO"]) + tokens(r["MATERNO"]))
        plan.append({"ID_PERSONA": r["ID_PERSONA"], "CODIGO": r["CODIGO"], "TIPO_DOCUMENTO": r["TIPO_DOCUMENTO"],
                     "ID_TIPODOCUMENTO": r["ID_TIPODOCUMENTO"], "NUM_DOCUMENTO": r["NUM_DOCUMENTO"],
                     "NOMBRES": r["NOMBRES"], "PATERNO": r["PATERNO"], "MATERNO": r["MATERNO"],
                     "SEDE": r["SEDE"], "PROGRAMA": r["PROGRAMA"], "CORREO_INST_ACTUAL_ORACLE": r["CORREO_INST_ACTUAL_ORACLE"],
                     "upn": f"{elegido}@{DOM}", "regla": regla, "usuario_lamb": propio, "displayName": nombre_vis,
                     "givenName": " ".join(x.capitalize() for x in tokens(r["NOMBRES"])),
                     "surname": " ".join(x.capitalize() for x in tokens(r["PATERNO"]) + tokens(r["MATERNO"]))})
    fecha = datetime.date.today().isoformat()
    pj = os.path.join(HOME, f"backups/entra-snapshots/plan_creacion_{fecha}.json")
    json.dump(plan, open(pj, "w"), ensure_ascii=False, indent=1); os.chmod(pj, 0o600)
    import openpyxl as ox
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = ox.Workbook(); ws = wb.active; ws.title = "Se crearían"
    cab = ["ID_PERSONA", "CODIGO", "NOMBRES", "PATERNO", "MATERNO", "SEDE", "PROGRAMA", "USUARIO_LAMB", "CUENTA_PROPUESTA", "ORIGEN_DEL_NOMBRE", "NOMBRE_VISIBLE"]
    ws.append(cab); [ws.append([p["ID_PERSONA"], p["CODIGO"], p["NOMBRES"], p["PATERNO"], p["MATERNO"], p["SEDE"], p["PROGRAMA"], p["usuario_lamb"], p["upn"], p["regla"], p["displayName"]]) for p in plan]
    ws2 = wb.create_sheet("Retenidos"); ws2.append(["ID_PERSONA", "CODIGO", "NOMBRES", "PATERNO", "MATERNO", "SEDE", "PROGRAMA", "MOTIVO"])
    for r, m in ret: ws2.append([r["ID_PERSONA"], r["CODIGO"], r["NOMBRES"], r["PATERNO"], r["MATERNO"], r["SEDE"], r["PROGRAMA"], m])
    for w in (ws, ws2):
        for c in w[1]: c.font, c.fill, c.alignment = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E79"), Alignment(wrap_text=True)
        for i in range(1, w.max_column + 1): w.column_dimensions[ox.utils.get_column_letter(i)].width = 22
        w.freeze_panes = "B2"; w.auto_filter.ref = w.dimensions
    xo = os.path.join(OUTD, f"Simulacion_creacion_cuentas_pregrado_{fecha}.xlsx"); wb.save(xo)
    from collections import Counter
    log(f"\nSE CREARÍAN: {len(plan)}  {dict(Counter(p['regla'] for p in plan))}")
    log(f"RETENIDOS:   {len(ret)}  {dict(Counter(re.sub(r'[«(:].*', '', m).strip() for _, m in ret))}")
    log(f"Excel: {xo}\nPlan para ejecutar: {pj}")

# ───────────────────────────── EJECUTAR ─────────────────────────────
def ejecutar(pj, limite):
    pw = envfile(os.path.join(HOME, ".secrets/lamb-cuentas.env")).get("LAMB_INITIAL_PASSWORD")
    if not pw: sys.exit("✗ falta LAMB_INITIAL_PASSWORD en ~/.secrets/lamb-cuentas.env")
    plan = json.load(open(pj))[: limite or None]
    stamp = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
    hechos, err = [], []
    for i, p in enumerate(plan, 1):
        # idempotencia: ¿ya existe una cuenta con este employeeId?
        st, j = graph("GET", "/users?$count=true&$select=id,userPrincipalName&$filter=" +
                      urllib.parse.quote(f"employeeId eq '{p['ID_PERSONA']}'"), eventual=True)
        if st == 200 and j["value"]:
            log(f"  ya existe ({j['value'][0]['userPrincipalName']}) → se salta {p['CODIGO']}")
            hechos.append({**p, "upn": j["value"][0]["userPrincipalName"], "estado": "ya existía"}); continue
        # la dirección debe seguir libre EN VIVO (UPN, mail y alias)
        a = p["upn"]
        st, j = graph("GET", "/users?$count=true&$select=id&$filter=" + urllib.parse.quote(
            f"userPrincipalName eq '{a}' or mail eq '{a}' or proxyAddresses/any(x:x eq 'smtp:{a}')"), eventual=True)
        if st != 200 or j["value"]:
            err.append((p["CODIGO"], a, "dirección ocupada o no verificable al ejecutar")); log(f"  ✗ {a} ocupada → no se crea"); continue
        body = {"accountEnabled": True, "displayName": p["displayName"], "givenName": p["givenName"], "surname": p["surname"],
                "mailNickname": a.split("@")[0], "userPrincipalName": a, "usageLocation": "PE",
                "employeeId": str(p["ID_PERSONA"]), "employeeType": "Estudiante", "companyName": "Universidad Peruana Unión",
                "passwordProfile": {"password": pw, "forceChangePasswordNextSignIn": True,
                                    "forceChangePasswordNextSignInWithMfa": True}}
        st, u = graph("POST", "/users", body)
        if st != 201:
            err.append((p["CODIGO"], a, f"crear: {st} {str(u)[:160]}")); log(f"  ✗ crear {a}: {st} {str(u)[:160]}"); continue
        log(f"  ✓ creada {a}  ({i}/{len(plan)})")
        hechos.append({**p, "id": u["id"], "upn": u["userPrincipalName"], "estado": "creada"})
    # licencia A3 con reintentos (usageLocation recién puesto tarda en propagarse)
    cola = [h for h in hechos if h.get("estado") == "creada"]
    if cola: log("esperando 60 s antes de licenciar…"); time.sleep(60)
    for ronda in range(1, 5):
        sig = []
        for h in cola:
            st, r = graph("POST", f"/users/{h['id']}/assignLicense", {"addLicenses": [{"skuId": A3, "disabledPlans": []}], "removeLicenses": []})
            if st == 200: h["licencia"] = "A3"
            elif st == 400 and "usage location" in str(r).lower(): sig.append(h)
            else: h["licencia"] = f"ERROR {st}"; err.append((h["CODIGO"], h["upn"], f"licencia: {st} {str(r)[:120]}"))
        if not sig: break
        log(f"ronda {ronda}: {len(sig)} pendientes de propagación · reintento en 45 s"); cola = sig; time.sleep(45)
    out = os.path.join(HOME, f"backups/entra-snapshots/creacion_resultado_{stamp}.json")
    json.dump({"hechos": hechos, "errores": err}, open(out, "w"), ensure_ascii=False, indent=1); os.chmod(out, 0o600)
    # Excel + SQL para Oracle con el UPN REALMENTE creado
    q = lambda s: "'" + str(s).replace("'", "''") + "'"
    creadas = [h for h in hechos if h.get("estado") in ("creada", "ya existía")]
    sql = [f"-- CORREO_INST de cuentas creadas en M365 · {stamp} · {len(creadas)} filas · sin COMMIT", "SET SERVEROUTPUT ON SIZE UNLIMITED", "SET DEFINE OFF",
           "DECLARE n NUMBER := 0; BEGIN"]
    for h in creadas:
        cur = (h.get("CORREO_INST_ACTUAL_ORACLE") or "").strip().lower()
        sql.append(f"  UPDATE MOISES.PERSONA_NATURAL SET CORREO_INST = {q(h['upn'].lower())} WHERE ID_PERSONA = {int(h['ID_PERSONA'])}"
                   f" AND ID_TIPODOCUMENTO = {int(h['ID_TIPODOCUMENTO'])} AND TRIM(NUM_DOCUMENTO) = {q(h['NUM_DOCUMENTO'])}"
                   f" AND NVL(LOWER(TRIM(CORREO_INST)),'#') = {q(cur or '#')}; n := n + SQL%ROWCOUNT;")
    sql += [f"  DBMS_OUTPUT.PUT_LINE('actualizadas: ' || n || ' de {len(creadas)}');", "END;", "/",
            f"-- Si 'actualizadas' = {len(creadas)} → COMMIT; si no → ROLLBACK;"]
    so = os.path.join(OUTD, f"UPDATE_CORREO_INST_cuentas_nuevas_{stamp}.sql"); open(so, "w").write("\n".join(sql) + "\n")
    log(f"\nCREADAS: {sum(1 for h in hechos if h.get('estado')=='creada')} · ya existían: {sum(1 for h in hechos if h.get('estado')=='ya existía')}"
        f" · con A3: {sum(1 for h in hechos if h.get('licencia')=='A3')} · errores: {len(err)}")
    log(f"resultado: {out}\nSQL para Oracle (Uziel): {so}")

if __name__ == "__main__":
    if len(sys.argv) < 3 or sys.argv[1] not in ("simular", "ejecutar"): sys.exit(__doc__)
    if sys.argv[1] == "simular": simular(sys.argv[2])
    else:
        lim = int(sys.argv[sys.argv.index("--limite") + 1]) if "--limite" in sys.argv else 0
        ejecutar(sys.argv[2], lim)
